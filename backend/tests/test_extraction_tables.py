from io import BytesIO

from app.services.chunking import chunk_text
from app.services.extraction import extract_text
from docx import Document


def test_docx_table_headers_and_rows_are_preserved_in_order():
    document = Document()
    document.add_paragraph("Procurement thresholds")
    table = document.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ("Amount", "Currency", "Approver")):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ("60,000", "INR", "Finance and Procurement")):
        cell.text = value
    content = BytesIO()
    document.save(content)

    extracted = extract_text("thresholds.docx", content.getvalue())
    assert extracted.index("Procurement thresholds") < extracted.index("Table 1:")
    assert "Table 1, row 1: Column 1: Amount | Column 2: Currency | Column 3: Approver" in extracted
    assert "Column 1: 60,000" in extracted
    assert "Column 2: INR" in extracted
    assert "Column 3: Finance and Procurement" in extracted
    assert all("60,000" in chunk for chunk in chunk_text(extracted, chunk_size=1000, overlap=0))


def test_empty_docx_table_does_not_count_as_extractable_text():
    document = Document()
    document.add_table(rows=1, cols=1)
    content = BytesIO()
    document.save(content)

    try:
        extract_text("empty.docx", content.getvalue())
    except ValueError as error:
        assert "No extractable text" in str(error)
    else:
        raise AssertionError("An empty DOCX should be rejected")


def test_pdf_text_extraction_keeps_page_provenance(monkeypatch):
    from app.services import extraction

    class Page:
        def extract_text(self):
            return "Quarterly total: INR 60,000"

    class Pdf:
        def __init__(self):
            self.pages = [Page()]

    monkeypatch.setattr(extraction, "PdfReader", lambda _: Pdf())
    monkeypatch.setattr(extraction, "_ocr_enabled", lambda: False)
    assert extraction.extract_text("report.pdf", b"pdf") == "[Page 1]\nQuarterly total: INR 60,000"


def test_scanned_page_uses_ocr_only_when_enabled(monkeypatch):
    from app.services import extraction

    class Page:
        def extract_text(self):
            return ""

    class Pdf:
        def __init__(self):
            self.pages = [Page()]

    monkeypatch.setattr(extraction, "PdfReader", lambda _: Pdf())
    monkeypatch.setattr(extraction, "_ocr_enabled", lambda: True)
    monkeypatch.setattr(extraction, "_ocr_pdf_page", lambda content, number: "Scanned amount: INR 500")
    assert "[Page 1]\nScanned amount: INR 500" in extraction.extract_text("scan.pdf", b"pdf")


def test_ocr_failure_is_reported_without_provider_exception_details(monkeypatch):
    import sys
    import types

    from app.services import extraction

    fake_pdf2image = types.SimpleNamespace(convert_from_bytes=lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("private renderer detail")))
    fake_tesseract = types.SimpleNamespace(image_to_string=lambda image: "")
    monkeypatch.setitem(sys.modules, "pdf2image", fake_pdf2image)
    monkeypatch.setitem(sys.modules, "pytesseract", fake_tesseract)
    try:
        extraction._ocr_pdf_page(b"private file bytes", 3)
    except ValueError as error:
        assert "page 3" in str(error)
        assert "private renderer detail" not in str(error)
    else:
        raise AssertionError("OCR errors should be converted to a safe extraction failure")
